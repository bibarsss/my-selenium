from multiprocessing import Process
from pathlib import Path
import sqlite3

from openpyxl import load_workbook
from flow_types.baseWithExcel import WithExcelType
import re

class MoveBy2SubfolderType(WithExcelType):
    @property
    def type(self):
        return self.__type

    @type.setter
    def type(self, value):
        self.__type = value

    def label(self)->str:
        return 'Переместить папки через 2 подпапки'

    def table_name(self)->str:
        return 'move_by_2_subfolder'

    def excel_map(self):
        return {}

    def save_to_excel(self):
        return

    def migration(self):
        connection = sqlite3.connect(self.cfg.get('db_name'))
        cursor = connection.cursor()
        cursor.execute(f'''
            DROP TABLE IF EXISTS {self.table_name()}
            ''')

        cursor.execute(f'''
            CREATE TABLE IF NOT EXISTS {self.table_name()}(
                        id INTEGER PRIMARY KEY,
                        group_id TEXT,
                        excel_line_number INTEGER,
                        search_text TEXT,
                        sud_folder_name TEXT,
                        client_folder_name TEXT,
                        status TEXT DEFAULT '',
                        status_text TEXT DEFAULT ''
                    )
                            ''')
        connection.commit()
        connection.close()

    def start(self):
        types = {
            1: 'Переместить папки через 2 подпапки',
            2: 'Выборка документов',
            3: 'Переместить папки через 2 подпапки (с группировкой)',
        }
        options = ".\n".join(f"{k} -> {v}" for k, v in types.items())
        try:
            print('==========================')
            print(options)
            print('==========================')

            self.type = int(input(f"Введите тип флоу: "))
            if self.type not in types.keys():
                print("Неправильный тип флоу: ", self.type)
                return
        except Exception as e:
            print("Неправильный тип флоу!")
            return

        super().start()

    def insert(self, row: tuple, cursor: sqlite3.Cursor, i):
        def safe_get(column_name: str) -> str:
            try:
                idx = self.cfg.index(column_name)
                return str(row[idx].value) if row[idx].value is not None else ""
            except (ValueError, IndexError):
                return ""

        data = {
            'search_text': safe_get('moveby2subfolder_excel_search_text'),
            'sud_folder_name': safe_get('moveby2subfolder_excel_sud_folder_name'),
            'client_folder_name': safe_get('moveby2subfolder_excel_client_folder_name'),
            "excel_line_number": i,
            "group_id": safe_get('moveby2subfolder_excel_group_id'),
            }

        columns = ", ".join(data.keys())
        placeholders = ", ".join([":" + key for key in data.keys()])
        query = f"INSERT INTO {self.table_name()}({columns}) VALUES ({placeholders})"

        cursor.execute(query, data)

    def _process_rows(self, ids, worker_id):
        print(f"[Worker {worker_id}] starting...")
        connection = sqlite3.connect(self.cfg.get('db_name'), timeout=30)
        connection.execute("PRAGMA journal_mode=WAL;")
        connection.execute("PRAGMA synchronous=NORMAL;")
        connection.execute("PRAGMA busy_timeout = 5000;")
        connection.row_factory = sqlite3.Row

        placeholder = ','.join('?' * len(ids))
        rows = connection.execute(f"SELECT * FROM {self.table_name()} WHERE id in ({placeholder})", ids).fetchall()

        for row in rows:
            excel_line_number = row['excel_line_number']
            print(f"[Worker {worker_id}] row: {excel_line_number}")
            self.run(row)

        connection.commit()
        connection.close()

    def run(self, row):
        def sanitize(name: str) -> str:
            return re.sub(r'[<>:"/\\|?*]', '', name.strip())

        main = sanitize(self.cfg.get('moveby2subfolder_main_folder_name'))
        sud = sanitize(row['sud_folder_name'])
        client = sanitize(row['client_folder_name'])
        grouped = sanitize(row['group_id'])

        if self.type == 1:
            target_dir = Path(main) / sud / client
        elif self.type == 2:
            target_dir = Path(main)
        elif self.type == 3:
            target_dir = Path(main) / grouped /client

        for file in Path(".").rglob("*"):
            if file.name.startswith("~$"):
                continue
            if file.suffix.lower() not in (".pdf", ".docx"):
                continue
            if target_dir in file.parents:
                continue
            try:
                if row['search_text'] in file.name:
                    target_file = target_dir / file.name
                    if target_file.exists():
                        file.unlink()
                    else:
                        target_dir.mkdir(parents=True, exist_ok=True)
                        file.rename(target_file)
            except Exception:
                continue
